from __future__ import annotations

import csv
import io
import re
from pathlib import Path

from parser.enrichment import enqueue_octopart_enrichment
from parser.pral.core.base import BaseParser
from parser.schema import BOMRow, ParsedOutput
from parser.pral.interfaces.parser_plugin import ParseContext, apply_capabilities

# Column header aliases → canonical field
_COLUMN_ALIASES: dict[str, list[str]] = {
    "mpn": [
        "mpn",
        "mfr part",
        "mfr part number",
        "manufacturer part number",
        "part number",
        "partnumber",
        "part #",
        "part",
        "pn",
        "p/n",
    ],
    "quantity": ["qty", "quantity", "count", "amount", "q'ty"],
    "ref": [
        "ref",
        "reference",
        "refdes",
        "ref des",
        "designator",
        "designators",
        "refs",
        "reference designator",
    ],
    "manufacturer": ["manufacturer", "mfr", "vendor", "supplier", "make"],
    "description": ["description", "desc", "value", "comment", "notes"],
}


class BOMParser(BaseParser):
    parser_name = "BOMParser"
    file_type = "BOM"

    def parse(self, file_bytes: bytes, context: ParseContext) -> ParsedOutput:
        output = self._base_output(context)
        output.parser_version = "1.0.0"
        ext = Path(context.filename).suffix.lower()

        try:
            if ext in (".xlsx", ".xls"):
                rows, headers = _parse_xlsx(file_bytes)
            else:
                rows, headers = _parse_csv(file_bytes)

            mapping = _normalize_headers(headers)
            if "mpn" not in mapping and "ref" not in mapping:
                output.warnings.append("missing required columns (MPN or Reference)")
                return apply_capabilities(output)

            bom_rows: list[BOMRow] = []
            mpns: set[str] = set()
            for row in rows:
                br = _row_to_bom(row, mapping)
                if br:
                    bom_rows.append(br)
                    if br.mpn:
                        mpns.add(br.mpn)

            output.bom_rows = bom_rows
            output.extracted_metadata["row_count"] = len(bom_rows)
            output.extracted_metadata["column_mapping"] = mapping

            if mpns:
                enqueue_octopart_enrichment(sorted(mpns), context.version_id)
        except Exception as exc:  # noqa: BLE001
            output.errors.append(f"parse error: {exc}")

        return apply_capabilities(output)


def _normalize_key(header: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", header.strip().lower()).strip()


def _normalize_headers(headers: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    normalized = [_normalize_key(h) for h in headers]
    for field, aliases in _COLUMN_ALIASES.items():
        for idx, nh in enumerate(normalized):
            if nh in aliases or any(a in nh for a in aliases):
                mapping[field] = idx
                break
    return mapping


def _parse_csv(file_bytes: bytes) -> tuple[list[list[str]], list[str]]:
    text = file_bytes.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return [], []
    return rows[1:], rows[0]


def _parse_xlsx(file_bytes: bytes) -> tuple[list[list[str]], list[str]]:
    from openpyxl import load_workbook  # noqa: PLC0415

    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        return [], []
    all_rows: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        all_rows.append([str(c) if c is not None else "" for c in row])
    wb.close()
    if not all_rows:
        return [], []
    return all_rows[1:], all_rows[0]


def _row_to_bom(row: list[str], mapping: dict[str, int]) -> BOMRow | None:
    def cell(field: str) -> str | None:
        idx = mapping.get(field)
        if idx is None or idx >= len(row):
            return None
        val = row[idx].strip()
        return val or None

    ref_raw = cell("ref")
    mpn = cell("mpn")
    if not ref_raw and not mpn:
        return None

    refs = re.split(r"[,;\s]+", ref_raw) if ref_raw else []
    refs = [r.strip() for r in refs if r.strip()]

    qty_raw = cell("quantity")
    qty = 1
    if qty_raw:
        try:
            qty = int(float(qty_raw))
        except ValueError:
            qty = 1

    return BOMRow(
        ref_designators=refs,
        mpn=mpn,
        manufacturer=cell("manufacturer"),
        quantity=qty,
        description=cell("description"),
    )
